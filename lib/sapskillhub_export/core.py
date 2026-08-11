from __future__ import annotations

import csv
import hashlib
import json
import re
import time
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping, Sequence

import openpyxl

SUPPORTED_OPTIONS = frozenset({"EQ", "NE", "BT", "NB", "GE", "GT", "LE", "LT", "CP", "NP"})
IDENTIFIER = re.compile(r"^[A-Z][A-Z0-9_]{0,29}$")
EXCEL_DATA_ROWS = 1_048_575


class SapGuiBusyTimeout(TimeoutError):
    """SAP GUI did not become idle before the bounded timeout."""


class ExportCompletionTimeout(TimeoutError):
    """An export file was missing, changing, locked, or SAP stayed busy."""


@dataclass(frozen=True)
class Filter:
    field: str
    sign: str = "I"
    option: str = "EQ"
    low: str = ""
    high: str = ""

    @classmethod
    def from_mapping(cls, value: Mapping[str, Any]) -> "Filter":
        item = cls(
            field=str(value.get("field", "")).strip().upper(),
            sign=str(value.get("sign", "I")).strip().upper(),
            option=str(value.get("option", "EQ")).strip().upper(),
            low=str(value.get("low", "")).strip(),
            high=str(value.get("high", "")).strip(),
        )
        item.validate()
        return item

    def validate(self) -> None:
        if not IDENTIFIER.fullmatch(self.field):
            raise ValueError(f"invalid SAP field name: {self.field!r}")
        if self.sign not in {"I", "E"}:
            raise ValueError(f"invalid sign for {self.field}: {self.sign!r}")
        if self.option not in SUPPORTED_OPTIONS:
            raise ValueError(f"unsupported option for {self.field}: {self.option!r}")
        if not self.low:
            raise ValueError(f"low is required for {self.field}")
        if self.option in {"BT", "NB"} and not self.high:
            raise ValueError(f"high is required for {self.field} {self.option}")


@dataclass(frozen=True)
class ChunkSpec:
    field: str
    type: str
    low: str
    high: str

    @classmethod
    def from_mapping(cls, value: Mapping[str, Any]) -> "ChunkSpec":
        result = cls(
            str(value.get("field", "")).strip().upper(),
            str(value.get("type", "")).strip().lower(),
            str(value.get("low", "")).strip(),
            str(value.get("high", "")).strip(),
        )
        if not IDENTIFIER.fullmatch(result.field):
            raise ValueError(f"invalid chunk field: {result.field!r}")
        if result.type not in {"integer", "padded_integer", "date"}:
            raise ValueError(f"unsupported chunk type: {result.type!r}")
        if not result.low or not result.high:
            raise ValueError("chunk low and high are required")
        if chunk_value(result.low, result.type) > chunk_value(result.high, result.type):
            raise ValueError("chunk low must not exceed high")
        return result


@dataclass(frozen=True)
class TablePolicy:
    required_fields: tuple[str, ...]
    chunk_field: str
    chunk_type: str
    key_fields: tuple[str, ...]


TABLE_POLICIES: dict[str, TablePolicy] = {
    "BSEG": TablePolicy(("BUKRS", "GJAHR"), "BELNR", "padded_integer", ("BUKRS", "GJAHR", "BELNR", "BUZEI")),
    "BSIK": TablePolicy(("BUKRS", "GJAHR"), "BELNR", "padded_integer", ("BUKRS", "GJAHR", "BELNR", "BUZEI")),
    "BSIS": TablePolicy(("BUKRS", "GJAHR"), "BELNR", "padded_integer", ("BUKRS", "GJAHR", "BELNR", "BUZEI")),
    "EKBE": TablePolicy(("BUDAT",), "BUDAT", "date", ("EBELN", "EBELP", "ZEKKN", "VGABE", "GJAHR", "BELNR", "BUZEI")),
}


@dataclass(frozen=True)
class Selection:
    schema_version: int
    filters: tuple[Filter, ...]
    columns: tuple[str, ...]
    sort: tuple[str, ...]
    chunk: ChunkSpec | None
    key_fields: tuple[str, ...]


@dataclass(frozen=True)
class ChunkRange:
    low: str
    high: str
    depth: int = 0


def normalize_identifiers(values: Iterable[Any], label: str) -> tuple[str, ...]:
    result = tuple(str(value).strip().upper() for value in values if str(value).strip())
    invalid = [value for value in result if not IDENTIFIER.fullmatch(value)]
    if invalid:
        raise ValueError(f"invalid {label}: {', '.join(invalid)}")
    return result


def parse_shorthand(expression: str, sign: str = "I") -> list[Filter]:
    if "=" not in expression:
        raise ValueError(f"filter must use FIELD=value: {expression!r}")
    field, raw = expression.split("=", 1)
    field = field.strip().upper()
    raw = raw.strip()
    if not raw:
        raise ValueError(f"filter value is empty: {expression!r}")
    if ".." in raw:
        low, high = raw.split("..", 1)
        return [Filter.from_mapping({"field": field, "sign": sign, "option": "BT", "low": low, "high": high})]
    return [Filter.from_mapping({"field": field, "sign": sign, "option": "EQ", "low": value}) for value in raw.split(",")]


def load_selection(path: Path | None = None, where: Sequence[str] = (), exclude: Sequence[str] = ()) -> Selection:
    payload: dict[str, Any] = {}
    if path:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(payload, dict):
            raise ValueError("selection JSON must contain an object")
    version = int(payload.get("schema_version", 1))
    if version != 1:
        raise ValueError(f"unsupported schema_version: {version}")
    filters = [Filter.from_mapping(value) for value in payload.get("filters", [])]
    for expression in where:
        filters.extend(parse_shorthand(expression, "I"))
    for expression in exclude:
        filters.extend(parse_shorthand(expression, "E"))
    chunk = ChunkSpec.from_mapping(payload["chunk"]) if payload.get("chunk") else None
    return Selection(
        version,
        tuple(filters),
        normalize_identifiers(payload.get("columns", []), "column names"),
        normalize_identifiers(payload.get("sort", []), "sort fields"),
        chunk,
        normalize_identifiers(payload.get("key_fields", []), "key fields"),
    )


def grouped_filter_semantics(filters: Sequence[Filter]) -> dict[str, dict[str, list[dict[str, str]]]]:
    grouped: dict[str, dict[str, list[dict[str, str]]]] = defaultdict(lambda: {"include_or": [], "exclude": []})
    for item in filters:
        target = "include_or" if item.sign == "I" else "exclude"
        grouped[item.field][target].append({"option": item.option, "low": item.low, "high": item.high})
    return dict(grouped)


def policy_for(table: str, selection: Selection) -> tuple[ChunkSpec | None, tuple[str, ...]]:
    table = table.upper()
    policy = TABLE_POLICIES.get(table)
    fields = {item.field for item in selection.filters}
    if policy:
        bounded_fields = {
            item.field
            for item in selection.filters
            if item.sign == "I" and item.option in {"EQ", "BT"}
        }
        missing = sorted(set(policy.required_fields) - bounded_fields)
        if missing:
            raise ValueError(
                f"{table} full export requires inclusive EQ/BT filter(s): {', '.join(missing)}"
            )
        chunk = selection.chunk
        if chunk is None:
            ranges = [
                item
                for item in selection.filters
                if item.field == policy.chunk_field
                and item.option == "BT"
                and item.sign == "I"
            ]
            if ranges:
                chunk = ChunkSpec(policy.chunk_field, policy.chunk_type, ranges[0].low, ranges[0].high)
        if chunk is None:
            raise ValueError(
                f"{table} full export requires a bounded {policy.chunk_field} range"
            )
        if chunk.field != policy.chunk_field or chunk.type != policy.chunk_type:
            raise ValueError(
                f"{table} requires chunk field/type {policy.chunk_field}/{policy.chunk_type}"
            )
        keys = selection.key_fields or policy.key_fields
        missing_keys = sorted(set(policy.key_fields) - set(keys))
        if missing_keys:
            raise ValueError(
                f"{table} key_fields must include: {', '.join(policy.key_fields)}"
            )
        validate_chunk_filter_compatibility(selection.filters, chunk)
        return chunk, keys
    if selection.chunk:
        validate_chunk_filter_compatibility(selection.filters, selection.chunk)
    return selection.chunk, selection.key_fields


def validate_chunk_filter_compatibility(
    filters: Sequence[Filter], spec: ChunkSpec
) -> None:
    includes = [
        item for item in filters if item.field == spec.field and item.sign == "I"
    ]
    if not includes:
        return
    if len(includes) != 1:
        raise ValueError(
            f"chunk field {spec.field} must have zero or one inclusive BT filter"
        )
    item = includes[0]
    if item.option != "BT" or item.low != spec.low or item.high != spec.high:
        raise ValueError(
            f"chunk field {spec.field} filter must exactly match chunk low/high"
        )


def chunk_value(value: str, value_type: str) -> int | date:
    if value_type in {"integer", "padded_integer"}:
        return int(value)
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"invalid chunk date: {value!r}")


def format_chunk_value(value: int | date, value_type: str, template: str) -> str:
    if value_type == "padded_integer":
        return str(value).zfill(len(template))
    if value_type == "integer":
        return str(value)
    if "-" in template:
        return value.strftime("%Y-%m-%d")  # type: ignore[union-attr]
    if "." in template:
        return value.strftime("%Y.%m.%d")  # type: ignore[union-attr]
    return value.strftime("%Y%m%d")  # type: ignore[union-attr]


def split_chunk(value: ChunkRange, spec: ChunkSpec) -> tuple[ChunkRange, ChunkRange]:
    low = chunk_value(value.low, spec.type)
    high = chunk_value(value.high, spec.type)
    if low >= high:
        raise RuntimeError(f"chunk {value.low}..{value.high} cannot be split further")
    if isinstance(low, date):
        midpoint = low + timedelta(days=(high - low).days // 2)
        next_value = midpoint + timedelta(days=1)
    else:
        midpoint = (low + high) // 2
        next_value = midpoint + 1
    return (
        ChunkRange(value.low, format_chunk_value(midpoint, spec.type, value.low), value.depth + 1),
        ChunkRange(format_chunk_value(next_value, spec.type, value.low), value.high, value.depth + 1),
    )


def recursive_chunks(spec: ChunkSpec, probe: Callable[[ChunkRange, int], int], chunk_size: int = 50_000, max_depth: int = 64) -> list[tuple[ChunkRange, int]]:
    leaves: list[tuple[ChunkRange, int]] = []

    def visit(current: ChunkRange) -> None:
        if current.depth > max_depth:
            raise RuntimeError("maximum chunk recursion depth exceeded")
        count = probe(current, chunk_size + 1)
        if count <= chunk_size:
            leaves.append((current, count))
            return
        left, right = split_chunk(current, spec)
        visit(left)
        visit(right)

    visit(ChunkRange(spec.low, spec.high))
    return leaves


def validate_chunk_coverage(chunks: Sequence[ChunkRange], spec: ChunkSpec) -> None:
    if not chunks:
        raise ValueError("no successful chunks")
    ordered = sorted(chunks, key=lambda item: chunk_value(item.low, spec.type))
    if ordered[0].low != spec.low or ordered[-1].high != spec.high:
        raise ValueError("chunk ranges do not cover the requested boundary")
    for previous, current in zip(ordered, ordered[1:]):
        left = chunk_value(previous.high, spec.type)
        right = chunk_value(current.low, spec.type)
        expected = left + (timedelta(days=1) if isinstance(left, date) else 1)
        if right != expected:
            raise ValueError(f"chunk gap or overlap between {previous.high} and {current.low}")


def read_workbook_matrix(path: Path) -> tuple[list[str], list[list[Any]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(values, ())]
        rows = [list(row) for row in values if any(value is not None for value in row)]
        return headers, rows
    finally:
        workbook.close()


def read_workbook_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    headers, matrix = read_workbook_matrix(path)
    if not headers or any(not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError(f"invalid or duplicate workbook headers: {path}")
    return headers, [dict(zip(headers, row)) for row in matrix]


def grid_column_ids(grid: Any) -> list[str]:
    value = getattr(grid, "ColumnOrder", ())
    columns: list[str] = []
    try:
        count = int(value.Count)
        for index in range(count):
            item = None
            for accessor in (
                lambda: value.Item(index),
                lambda: value.ElementAt(index),
                lambda: value(index),
            ):
                try:
                    item = accessor()
                    break
                except Exception:
                    continue
            if item is None:
                raise RuntimeError(f"cannot read ALV column at index {index}")
            columns.append(str(item))
    except Exception:
        try:
            columns = [str(item) for item in value]
        except Exception as exc:
            raise RuntimeError("cannot read ALV technical column IDs") from exc
    if not columns or any(not item for item in columns) or len(set(columns)) != len(columns):
        raise RuntimeError("ALV technical column IDs are empty or duplicated")
    return columns


def matrix_by_technical_columns(
    header_count: int,
    rows: Sequence[Sequence[Any]],
    technical_columns: Sequence[str],
) -> list[dict[str, Any]]:
    if header_count != len(technical_columns):
        raise ValueError(
            "ALV/workbook column count mismatch: "
            f"technical={len(technical_columns)}, workbook={header_count}"
        )
    return [
        {
            technical_columns[index]: row[index] if index < len(row) else None
            for index in range(header_count)
        }
        for row in rows
    ]


def rewrite_workbook_with_technical_headers(
    path: Path, technical_columns: Sequence[str]
) -> int:
    headers, rows = read_workbook_matrix(path)
    technical_rows = matrix_by_technical_columns(
        len(headers), rows, technical_columns
    )
    workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    sheet.append(list(technical_columns))
    for row in technical_rows:
        sheet.append([row.get(column) for column in technical_columns])
    workbook.save(path)
    return len(technical_rows)


def merge_workbooks(parts: Sequence[Path], xlsx_path: Path, csv_path: Path, key_fields: Sequence[str], sort_fields: Sequence[str] = (), overwrite: bool = False, output_columns: Sequence[str] = ()) -> dict[str, Any]:
    if not parts:
        raise ValueError("no part workbooks to merge")
    for output in (xlsx_path, csv_path):
        if output.exists() and not overwrite:
            raise FileExistsError(f"output exists; pass --overwrite: {output}")
    headers: list[str] | None = None
    rows: list[dict[str, Any]] = []
    for part in parts:
        current_headers, current_rows = read_workbook_rows(part)
        if headers is None:
            headers = current_headers
        elif current_headers != headers:
            raise ValueError(f"column structure mismatch: {part}")
        rows.extend(current_rows)
    headers = headers or []
    if output_columns:
        missing_columns = [field for field in output_columns if field not in headers]
        if missing_columns:
            raise ValueError(f"requested output column(s) missing: {', '.join(missing_columns)}")
        headers = list(output_columns)
    missing_keys = [field for field in key_fields if field not in headers]
    if missing_keys:
        raise ValueError(f"key field(s) missing from export: {', '.join(missing_keys)}")
    seen: set[tuple[str, ...]] = set()
    duplicates: list[tuple[str, ...]] = []
    if key_fields:
        for row in rows:
            key = tuple("" if row.get(field) is None else str(row.get(field)) for field in key_fields)
            if any(not value.strip() for value in key):
                raise ValueError(f"blank business key detected: {key}")
            if key in seen:
                duplicates.append(key)
            seen.add(key)
    if duplicates:
        raise ValueError(f"duplicate business key detected: {duplicates[0]}")
    order = tuple(sort_fields) or tuple(key_fields)
    rows.sort(key=lambda row: tuple("" if row.get(field) is None else str(row.get(field)) for field in order))
    rows = [{header: row.get(header) for header in headers} for row in rows]
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook(write_only=True)
    for offset in range(0, max(len(rows), 1), EXCEL_DATA_ROWS):
        sheet = workbook.create_sheet(f"Data-{offset // EXCEL_DATA_ROWS + 1:04d}")
        sheet.append(headers)
        for row in rows[offset : offset + EXCEL_DATA_ROWS]:
            sheet.append([row.get(header) for header in headers])
    workbook.save(xlsx_path)
    with csv_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    return {"columns": headers, "rows": len(rows), "duplicates": 0, "sheets": len(workbook.worksheets)}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_manifest(path: Path, payload: Mapping[str, Any], overwrite: bool = False) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"manifest exists; pass --overwrite: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def session_identity(session: Any) -> dict[str, str]:
    info = session.Info
    return {"system": str(getattr(info, "SystemName", "")), "client": str(getattr(info, "Client", ""))}


def connect_sap_session(expected_system: str = "", expected_client: str = "") -> Any:
    if __import__("os").name != "nt":
        raise RuntimeError("SAP GUI automation requires Windows")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("install pywin32 to use SAP GUI automation") from exc
    pythoncom.CoInitialize()
    application = win32com.client.GetObject("SAPGUI").GetScriptingEngine
    sessions: list[Any] = []
    for connection_index in range(application.Children.Count):
        connection = application.Children(connection_index)
        for session_index in range(connection.Children.Count):
            sessions.append(connection.Children(session_index))
    if not sessions:
        raise RuntimeError("no authenticated SAP GUI session is available")
    matches: list[Any] = []
    for session in sessions:
        identity = session_identity(session)
        if not identity["system"] or not identity["client"]:
            continue
        if expected_system and identity["system"].upper() != expected_system.upper():
            continue
        if expected_client and identity["client"] != expected_client:
            continue
        matches.append(session)
    if not matches:
        requested = f"system={expected_system or '*'}, client={expected_client or '*'}"
        available = ", ".join(f"{item['system']}/{item['client']}" for item in map(session_identity, sessions))
        raise RuntimeError(f"no SAP GUI session matches {requested}; available={available}")
    if len(matches) > 1:
        raise RuntimeError("multiple SAP GUI sessions match; specify --system and --client and leave only one matching session")
    return matches[0]


def find_required(session: Any, control_id: str) -> Any:
    try:
        return session.FindById(control_id)
    except Exception as exc:
        raise RuntimeError(f"required SAP technical control is missing: {control_id}") from exc


def set_text_verified(session: Any, control_id: str, value: str) -> None:
    control = find_required(session, control_id)
    control.Text = str(value)
    accepted = str(getattr(control, "Text", "")).strip()
    if accepted != str(value).strip():
        raise RuntimeError(f"SAP rejected value for {control_id}: requested={value!r}, accepted={accepted!r}")


def set_key_verified(session: Any, control_id: str, value: str) -> None:
    control = find_required(session, control_id)
    control.Key = str(value)
    accepted = str(getattr(control, "Key", "")).strip()
    if accepted != str(value).strip():
        raise RuntimeError(f"SAP rejected key for {control_id}: requested={value!r}, accepted={accepted!r}")


def load_control_profile(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if payload.get("schema_version") != 1:
        raise ValueError(f"unsupported control profile: {path}")
    return payload


def wait_ready(session: Any, timeout: float = 300) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if not bool(getattr(session, "Busy", False)):
            return
        time.sleep(0.25)
    raise SapGuiBusyTimeout("SAP GUI remained busy")


def wait_export_complete(
    session: Any,
    output: Path,
    *,
    timeout: float = 120,
    stable_for: float = 1.0,
    poll_interval: float = 0.25,
    clock: Callable[[], float] = time.monotonic,
    sleeper: Callable[[float], None] = time.sleep,
) -> None:
    """Wait until a non-empty file is stable, readable, and SAP is idle.

    File creation alone is insufficient because SAP GUI can create the XLSX
    before the XXL writer and its COM call have completed.
    """

    deadline = clock() + timeout
    last_signature: tuple[int, int] | None = None
    stable_since: float | None = None
    observed = "file was not created"
    while clock() < deadline:
        now = clock()
        signature: tuple[int, int] | None = None
        try:
            stat = output.stat()
            if stat.st_size > 0:
                signature = (stat.st_size, stat.st_mtime_ns)
                with output.open("rb") as stream:
                    stream.read(1)
        except (FileNotFoundError, OSError):
            signature = None
        if signature is None:
            observed = "file is missing, empty, or locked"
            last_signature = None
            stable_since = None
        elif signature != last_signature:
            observed = "file is still changing"
            last_signature = signature
            stable_since = now
        elif stable_since is not None and now - stable_since >= stable_for:
            if not bool(getattr(session, "Busy", False)):
                return
            observed = "file is stable but SAP GUI remains busy"
        sleeper(poll_interval)
    raise ExportCompletionTimeout(
        f"SAP export did not complete for {output}: {observed}"
    )


def status_error(session: Any) -> str:
    try:
        bar = session.FindById("wnd[0]/sbar")
        if str(getattr(bar, "MessageType", "")).upper() in {"E", "A"}:
            return str(getattr(bar, "Text", ""))
    except Exception:
        pass
    return ""


def export_alv(
    session: Any,
    grid_id: str,
    output: Path,
    overwrite: bool = False,
    *,
    completion_timeout: float = 120,
    stable_for: float = 1.0,
) -> None:
    if output.exists() and not overwrite:
        raise FileExistsError(f"output exists; pass --overwrite: {output}")
    if output.exists():
        output.unlink()
    output.parent.mkdir(parents=True, exist_ok=True)
    grid = find_required(session, grid_id)
    grid.PressToolbarContextButton("&MB_EXPORT")
    grid.SelectContextMenuItem("&XXL")
    find_required(session, "wnd[1]/tbar[0]/btn[20]").Press()
    set_text_verified(session, "wnd[1]/usr/ctxtDY_PATH", str(output.parent))
    set_text_verified(session, "wnd[1]/usr/ctxtDY_FILENAME", output.name)
    find_required(session, "wnd[1]/tbar[0]/btn[0]").Press()
    wait_export_complete(
        session,
        output,
        timeout=completion_timeout,
        stable_for=stable_for,
    )
